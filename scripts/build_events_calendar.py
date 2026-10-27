#!/usr/bin/env python3
"""Builds META_Events_Calendar.xlsx from events_calendar.json (port of build-events-calendar.ps1, no Excel needed).

  META_Intel/events_calendar.json  ->  META_Intel/META_Events_Calendar.xlsx

The JSON is the source of truth; the daily META Intel run edits the JSON and then runs this script.
Needs openpyxl (pip install openpyxl). Prints the saved path on the last line.
"""
import argparse
import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY, GREY, WHITE = "1F3864", "595959", "FFFFFF"
GREEN = {"fill": "C6EFCE", "font": "006100"}
AMBER = {"fill": "FFEB9C", "font": "9C5700"}
RED = {"fill": "FFC7CE", "font": "9C0006"}
SILVER = {"fill": "D9D9D9", "font": "404040"}
TOP = Alignment(vertical="top", wrap_text=True)


def font(**kw):
    return Font(name="Arial", size=kw.pop("size", 10), **kw)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def put(ws, row, col, value, **f):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font(**f)
    return c


def as_date(s):
    try:
        return dt.datetime.strptime(str(s), "%Y-%m-%d")
    except (TypeError, ValueError):
        return None


def add_table(ws, name, ref, style):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=str(Path(__file__).resolve().parent.parent / "META_Intel"))
    root = Path(ap.parse_args().root)
    data = json.loads((root / "events_calendar.json").read_text(encoding="utf-8-sig"))
    out_path = root / "META_Events_Calendar.xlsx"

    # Dated events first (by start date), undated (TBD / unverified) last.
    events = sorted(data.get("events") or [], key=lambda e: (0 if e.get("start") else 1, e.get("start") or "", e.get("event") or ""))

    wb = Workbook()
    ws = wb.active
    ws.title = "Calendar"
    wo = wb.create_sheet("Overview")
    wc = wb.create_sheet("Concluded")
    wl = wb.create_sheet("Change log")
    for sh in (ws, wo, wc, wl):
        sh.sheet_view.showGridLines = False

    # ---------------- Calendar ----------------
    put(ws, 1, 1, "META Events & Conferences Calendar", size=16, bold=True, color=NAVY)
    put(ws, 2, 1, f"Lenovo RHQ Riyadh - Market Intelligence  |  Last updated: {data.get('last_updated', '')}  |  Status colours and summary: see the Overview sheet", color=GREY)
    put(ws, 3, 1, f"Risk note: {data.get('risk_note', '')}", italic=True, color=RED["font"])
    ws.merge_cells("A3:H3")
    ws["A3"].alignment = TOP
    ws.row_dimensions[3].height = 40

    headers = ["Event", "Start date", "End date", "Days to start", "Dates (as reported)", "Status", "Region", "Country",
               "City / venue", "Lenovo BU relevance", "Key accounts / competitors", "Action", "Source + verified on"]
    widths = [34, 12, 12, 9, 24, 22, 9, 13, 30, 26, 42, 42, 38]
    hdr = 5
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        put(ws, hdr, i, h, bold=True, color=WHITE)
        ws.column_dimensions[get_column_letter(i)].width = w

    r = hdr + 1
    for e in events:
        put(ws, r, 1, str(e.get("event", "")), bold=True)
        for col, key in ((2, "start"), (3, "end")):
            d = as_date(e.get(key))
            c = put(ws, r, col, d)
            c.number_format = "dd-mmm-yyyy"
        put(ws, r, 4, f'=IF(B{r}="","",B{r}-TODAY())').number_format = "0"
        vals = [e.get(k, "") for k in ("dates_text", "status", "region", "country", "city_venue", "bu", "accounts", "action", "source")]
        for i, v in enumerate(vals):
            put(ws, r, 5 + i, "" if v is None else str(v))
        r += 1
    last = max(r - 1, hdr + 1)
    add_table(ws, "EventsCalendar", f"A{hdr}:M{last}", "TableStyleMedium2")
    for row in ws.iter_rows(min_row=hdr + 1, max_row=last, max_col=13):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if c.column in (2, 3, 4) else None)
    ws.row_dimensions[hdr].height = 30
    for c in ws[hdr]:
        c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center" if c.column in (2, 3, 4) else None)

    # Status colours.
    f = f"$F{hdr + 1}"
    rules = [
        (f'ISNUMBER(SEARCH("Postponed",{f}))', RED),
        (f'OR(ISNUMBER(SEARCH("risk",{f})),ISNUMBER(SEARCH("watch",{f})))', AMBER),
        (f'ISNUMBER(SEARCH("unconfirmed",{f}))', SILVER),
        (f'AND(ISNUMBER(SEARCH("Upcoming",{f})),ISERROR(SEARCH("risk",{f})),ISERROR(SEARCH("watch",{f})))', GREEN),
    ]
    for formula, c in rules:
        ws.conditional_formatting.add(f"F{hdr + 1}:F{last}", FormulaRule(formula=[formula], fill=fill(c["fill"]), font=Font(color=c["font"], bold=True), stopIfTrue=True))
    # Events starting within 7 days: bold red countdown.
    d1 = f"$D{hdr + 1}"
    ws.conditional_formatting.add(f"D{hdr + 1}:D{last}", FormulaRule(formula=[f"AND(ISNUMBER({d1}),{d1}>=0,{d1}<=7)"], fill=fill(RED["fill"]), font=Font(color=RED["font"], bold=True)))

    ws.freeze_panes = f"B{hdr + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hdr}:{hdr}"

    # ---------------- Concluded ----------------
    put(wc, 1, 1, "Concluded events (last 60 days)", size=14, bold=True, color=NAVY)
    for i, (h, w) in enumerate(zip(["Event", "Dates", "City / venue", "Outcome", "Source + verified on"], [36, 24, 28, 60, 40]), start=1):
        put(wc, 3, i, h, bold=True, color=WHITE)
        wc.column_dimensions[get_column_letter(i)].width = w
    r = 4
    for e in data.get("concluded") or []:
        for i, k in enumerate(("event", "dates_text", "city_venue", "outcome", "source"), start=1):
            c = put(wc, r, i, str(e.get(k) or ""), bold=(i == 1))
            c.alignment = TOP
        r += 1
    add_table(wc, "ConcludedEvents", f"A3:E{max(4, r - 1)}", "TableStyleMedium4")

    # ---------------- Change log ----------------
    put(wl, 1, 1, "Change log", size=14, bold=True, color=NAVY)
    put(wl, 3, 1, "Date", bold=True, color=WHITE)
    put(wl, 3, 2, "Change", bold=True, color=WHITE)
    wl.column_dimensions["A"].width = 13
    wl.column_dimensions["B"].width = 110
    r = 4
    for l in data.get("changelog") or []:
        c = put(wl, r, 1, as_date(l.get("date")) or l.get("date"))
        c.number_format = "dd-mmm-yyyy"
        c.alignment = Alignment(horizontal="left", vertical="top")
        put(wl, r, 2, str(l.get("change") or "")).alignment = TOP
        r += 1
    add_table(wl, "ChangeLog", f"A3:B{max(4, r - 1)}", "TableStyleLight9")

    # ---------------- Overview ----------------
    wo.column_dimensions["A"].width = 34
    wo.column_dimensions["B"].width = 14
    wo.column_dimensions["C"].width = 70
    put(wo, 1, 1, "Overview", size=14, bold=True, color=NAVY)
    D = f"Calendar!$D${hdr + 1}:$D${last}"
    F = f"Calendar!$F${hdr + 1}:$F${last}"
    G = f"Calendar!$G${hdr + 1}:$G${last}"
    summary = [
        ("Events tracked", f"=COUNTA(Calendar!$A${hdr + 1}:$A${last})"),
        ("Starting in the next 7 days", f'=COUNTIFS({D},">=0",{D},"<=7")'),
        ("Starting in the next 30 days", f'=COUNTIFS({D},">=0",{D},"<=30")'),
        ("At risk or on watch", f'=COUNTIF({F},"*risk*")+COUNTIF({F},"*watch*")'),
        ("Postponed", f'=COUNTIF({F},"*Postponed*")'),
        ("Dates unconfirmed", f'=COUNTIF({F},"*unconfirmed*")'),
        ("KSA events", f'=COUNTIF({G},"KSA")'),
        ("GCC events", f'=COUNTIF({G},"GCC")'),
        ("MEA events", f'=COUNTIF({G},"MEA")'),
        ("Global events", f'=COUNTIF({G},"Global")'),
    ]
    for col, h in ((1, "Summary"), (2, "Count")):
        c = put(wo, 3, col, h, bold=True, color=WHITE)
        c.fill = fill(NAVY)
    thin = Border(bottom=Side(style="thin", color="D9D9D9"))
    r = 4
    for label, formula in summary:
        put(wo, r, 1, label).border = thin
        c = put(wo, r, 2, formula)
        c.alignment = Alignment(horizontal="center")
        c.border = thin
        r += 1

    r += 1
    put(wo, r, 1, "Status colours", bold=True)
    r += 1
    legend = [
        ("Upcoming", GREEN, "Confirmed on the official site or reputable press; no known risk"),
        ("At risk / watch", AMBER, "Still scheduled, but security, venue or attendance risk; check daily"),
        ("Postponed", RED, "Moved or cancelled; new dates shown where announced"),
        ("Dates unconfirmed", SILVER, "No official confirmation yet; verify before planning"),
        ("Days to start, 0-7", RED, "Event starts within a week (countdown highlighted on the Calendar sheet)"),
    ]
    for label, c, desc in legend:
        cell = put(wo, r, 1, label, bold=True, color=c["font"])
        cell.fill = fill(c["fill"])
        put(wo, r, 3, desc)
        r += 1
    r += 1
    put(wo, r, 1, "How this file is maintained", bold=True)
    r += 1
    notes = [
        "Updated by the daily META Intel cloud run (07:00 KSA) from the events and conferences agent (meta-events-intel).",
        "Source of truth: events_calendar.json in the same folder. This workbook is rebuilt from it by scripts/build_events_calendar.py, so edits made only in Excel are overwritten on the next run.",
        "Every date carries its source and the date it was verified. 'Days to start' recalculates from today's date each time the file is opened.",
        f"Last updated: {data.get('last_updated', '')}.",
    ]
    for n in notes:
        put(wo, r, 1, n).alignment = TOP
        wo.merge_cells(f"A{r}:C{r}")
        wo.row_dimensions[r].height = 27
        r += 1

    wb.active = 0
    wb.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
